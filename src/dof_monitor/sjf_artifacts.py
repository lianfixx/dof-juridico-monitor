"""Generate four reviewable SJF files only after complete structured audit.

Outputs are NOT final until PDF/DOCX visual review and hash verification.
This module does not send email or certify that source documents were read.
"""
import argparse
import hashlib
import html
import json
from pathlib import Path
from urllib.parse import urlsplit

from .sjf_weekly import audit
from .sjf_intake import _load, _save

OPINION_SLOT = "Comentario personal: [Espacio para mi opinión antes del envío]"

def _words(value):
    return len(value.split())

def validate_editorial(data):
    errors = []
    if not isinstance(data, dict):
        return ["Editorial JSON missing"]
    summary = data.get("summary", "")
    if not isinstance(summary, str) or not summary.strip():
        errors.append("Lo relevante missing")
    items = data.get("items")
    if not isinstance(items, list) or not items:
        errors.append("No selected legal findings")
        items = []
    for n, item in enumerate(items, 1):
        for field in ("title", "effect", "status", "importance", "study", "source_url"):
            if not isinstance(item.get(field), str) or not item[field].strip():
                errors.append(f"Item {n}: {field} missing")
        url = item.get("source_url", "")
        host = urlsplit(url).hostname if isinstance(url, str) else None
        if not host or not (host == "scjn.gob.mx" or host.endswith(".scjn.gob.mx")):
            errors.append(f"Item {n}: non-official SCJN source")
    suggestion = data.get("suggested_comment", "")
    if not isinstance(suggestion, str) or not 40 <= _words(suggestion) <= 70:
        errors.append("Suggested comment must contain 40-70 words")
    email = data.get("draft_email", "")
    if not isinstance(email, str) or not 80 <= _words(email) <= 130:
        errors.append("BLUF draft must contain 80-130 words")
    if not isinstance(email, str) or OPINION_SLOT not in email:
        errors.append("Personal opinion placeholder missing")
    if any("{{" in str(v) or "}}" in str(v) for v in [summary, suggestion, email]):
        errors.append("Unresolved placeholders")
    return errors

def _hash(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()

def _pdf(path, title, summary, items, guide=False):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BodyCustom", parent=styles["Normal"], fontName="Helvetica",
                              fontSize=9, leading=13, spaceAfter=8, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="CellCustom", parent=styles["Normal"], fontName="Helvetica",
                              fontSize=8, leading=11))
    P = lambda s, style="BodyCustom": Paragraph(html.escape(str(s)).replace("\n","<br/>"), styles[style])
    story = [P(title, "Title"), Spacer(1, 12)]
    if guide:
        story.extend([P("Guía de estudio y control", "Heading2"), P(summary)])
        for item in items:
            story.extend([P(item["title"], "Heading2"), P(item["study"]),
                          P("Fuente oficial: " + item["source_url"])])
    else:
        story.extend([P("Lo relevante", "Heading2"), P(summary)])
        rows = [[P(s,"CellCustom") for s in
                 ("Asunto", "Qué ocurrió/efecto", "Vigencia/estado", "Por qué importa/qué revisar")]]
        for item in items:
            rows.append([P(item["title"],"CellCustom"),
                         P(item["effect"] + "\nFuente: " + item["source_url"],"CellCustom"),
                         P(item["status"],"CellCustom"), P(item["importance"],"CellCustom")])
        table = LongTable(rows, colWidths=[92,170,100,130], repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([("GRID",(0,0),(-1,-1),0.35,colors.grey),
                                   ("BACKGROUND",(0,0),(-1,0),colors.whitesmoke),
                                   ("VALIGN",(0,0),(-1,-1),"TOP"),
                                   ("LEFTPADDING",(0,0),(-1,-1),6),
                                   ("RIGHTPADDING",(0,0),(-1,-1),6)]))
        story.append(table)
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=44, rightMargin=44,
                            topMargin=42, bottomMargin=42)
    doc.build(story)

def _docx(path, title, summary, items, guide=False):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    section = doc.sections[0]
    section.header.paragraphs[0].text = "ACTUALIZACIÓN JURÍDICA"
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    doc.add_heading(title, 0)
    doc.add_heading("Guía de estudio" if guide else "Lo relevante", 1)
    doc.add_paragraph(summary)
    if guide:
        for item in items:
            doc.add_heading(item["title"], 2)
            doc.add_paragraph(item["study"])
            doc.add_paragraph("Fuente oficial: " + item["source_url"])
    else:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, value in zip(table.rows[0].cells,
            ("Asunto","Qué ocurrió/efecto","Vigencia/estado","Por qué importa/qué revisar")):
            cell.text = value
        for item in items:
            cells = table.add_row().cells
            for cell, value in zip(cells, (item["title"],
                item["effect"] + "\nFuente: " + item["source_url"],
                item["status"], item["importance"])):
                cell.text = value
    doc.save(path)

def build(control_path, editorial_path, output_dir):
    problems = audit(control_path)
    data = json.loads(Path(editorial_path).read_text(encoding="utf-8"))
    problems.extend(validate_editorial(data))
    if problems:
        raise ValueError("NO FINAL:\n" + "\n".join(problems))
    control = _load(control_path)
    edition = control["edition_date"]
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)
    report_title = "Semanario Judicial | " + edition
    guide_title = "Guía de estudio SJF | " + edition
    items = data["items"]
    outputs = {}
    for kind, guide in (("Reporte", False), ("Guia", True)):
        for extension in ("pdf", "docx"):
            path = out / f"SJF_{edition}_{kind}.{extension}"
            if path.exists():
                raise FileExistsError(f"Refusing to overwrite: {path}")
            if extension == "pdf":
                _pdf(path, guide_title if guide else report_title, data["summary"], items, guide)
            else:
                _docx(path, guide_title if guide else report_title, data["summary"], items, guide)
            outputs[path.name] = {"sha256": _hash(path), "render_checked": False}
    control["deliverables"] = outputs
    control["status"] = "READY_FOR_VISUAL_REVIEW"
    control["email"]["state"] = "NOT_SENT"
    _save(control_path, control)
    return outputs

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--control", required=True)
    p.add_argument("--editorial", required=True)
    p.add_argument("--out", required=True)
    args = p.parse_args()
    print(json.dumps(build(args.control,args.editorial,args.out),indent=2,ensure_ascii=False))
    print("NOT FINAL: render and inspect all four files, verify hashes and approve explicitly.")

if __name__ == "__main__":
    main()
